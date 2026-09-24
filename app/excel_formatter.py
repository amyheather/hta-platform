from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# =============================================================================
# Styles
# =============================================================================

HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9D9D9")

HEADER_FONT = Font(bold=True)

TITLE_FONT = Font(bold=True, size=13)

THIN = Side(style="thin", color="C0C0C0")

BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_ALIGN = Alignment(
    horizontal="center", vertical="center", wrap_text=True
)

BODY_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)

# =============================================================================
# Basic Formatting
# =============================================================================


def format_header(ws, row):
    """Format the header row."""

    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = HEADER_ALIGN


def format_body(ws, start_row=2):
    """Format worksheet body."""

    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = BODY_ALIGN


def auto_width(ws):
    """Automatically resize columns."""

    for column in ws.columns:

        width = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in column
        )

        ws.column_dimensions[get_column_letter(column[0].column)].width = min(
            width + 3, 50
        )


def add_filter(ws, header_row):
    """Apply Excel filter."""

    ws.auto_filter.ref = (
        f"A{header_row}:" f"{get_column_letter(ws.max_column)}" f"{ws.max_row}"
    )


# =============================================================================
# Generic Table Formatter
# =============================================================================


def format_standard_table(ws, header_row=1, body_start_row=2):
    """
    Generic formatter for any tabular worksheet.

    Suitable for:
    - Repository index
    - Web scraper output
    - Future Excel exports
    """

    format_header(ws, header_row)

    format_body(ws, start_row=body_start_row)

    auto_width(ws)

    add_filter(ws, header_row)


# =============================================================================
# Single Table Formatter
# =============================================================================


def format_single_table_excel(ws, metadata_rows):
    """
    Formatter for extracted table exports.
    """

    header_row = metadata_rows + 4

    ws["A1"].font = TITLE_FONT

    format_header(ws, header_row)

    format_body(ws, start_row=1)

    auto_width(ws)

    add_filter(ws, header_row)


# =============================================================================
# Repository Formatter
# =============================================================================


def format_repository_excel(ws):
    """
    Backwards-compatible wrapper.
    """

    format_standard_table(ws)
